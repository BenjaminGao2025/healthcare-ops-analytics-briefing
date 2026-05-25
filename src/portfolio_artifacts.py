"""Generate executive portfolio artifacts for the healthcare ops demo."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.dashboard_data import dashboard_summary, load_dashboard_frames


PROJECT_ROOT = Path(__file__).resolve().parents[1]
METHODOLOGY_FILE = Path("reports/methodology_note.md")
REVIEW_GUIDE_FILE = Path("reports/opus_review_guide.md")
DECK_SOURCE_FILE = Path("reports/board_briefing_deck.md")
DECK_PDF_FILE = Path("reports/board_briefing_deck.pdf")
EXCEL_FILE = Path("excel/healthcare_ops_analytics_workbook.xlsx")


def build_methodology_note() -> str:
    """Build a one-page methodology note."""
    return """# Methodology Note

## Purpose

This portfolio project demonstrates how public aggregate healthcare wait-time data can be transformed into executive-ready reporting, operational drilldowns, and data quality monitoring.

## Data Boundary

- Uses public aggregate sources only: CIHI wait-time tables and BC_MoH surgical wait-time exports.
- No patient-level data.
- No VCH internal data.
- No clinical decision support or patient-specific inference.

## Transformations

- CIHI national, provincial, and regional rows are normalized into shared procedure, geography, year, period, and KPI fields.
- BC_MoH fiscal-year surgical rows are normalized into province, health authority, and hospital geographies.
- BC_MoH percentile fields (`COMPLETED_50TH_PERCENTILE`, `COMPLETED_90TH_PERCENTILE`) are interpreted as weeks in the public export and converted to days; range checks against CIHI hip/knee procedure magnitudes are documented in `INSPECTION_NOTES.md`.
- CIHI Hip Fracture Repair values reported in hours are converted to days for table consistency and called out in the Executive Summary.
- Suppressed or unavailable case-volume values are preserved as nulls rather than estimated.

## Interpretation Limits

- CIHI and BC_MoH are not perfect substitutes; they answer overlapping but different reporting questions.
- Missing benchmark percentages and suppressed volumes should be visible in dashboard caveats.
- Load recency reflects when the database was refreshed, not the source publication date.
- Results are appropriate for portfolio demonstration and planning-style analytics, not for operational claims about live health-system performance.
"""


def build_review_guide() -> str:
    """Build the Opus review guide for a whole-project review."""
    return """# Opus Review Guide

## Review Goal

Review this repository as a portfolio project for a Junior Business Advisor, Data & Analytics application. Prioritize correctness, honesty, healthcare reporting fit, and whether the project demonstrates executive-ready communication.

## Architecture review

- Check whether the flow from raw public files to PostgreSQL to KPI/data-quality/EDA/reporting artifacts is clear and repeatable.
- Check whether module boundaries are reasonable: ingest, KPI, quality, EDA, dashboard data, and portfolio artifacts.
- Check whether generated reports match the stated public-data boundaries.

## Data review

- Verify that no patient-level or internal VCH data is used.
- Review the CIHI and BC_MoH unit assumptions, especially BC weeks-to-days conversion and CIHI hip-fracture hours-to-days conversion.
- Review the handling of suppressed/null case volumes and missing benchmark values.

## Dashboard review

- Assess whether the dashboard pages are useful for healthcare leaders and analysts.
- Check whether caveats are visible without overwhelming the executive summary.
- Review chart choices for clarity and misleading comparisons.

## Code review

- Run `make load`, `make kpis`, `make quality`, `make eda`, `make artifacts`, `pytest -q`, and `ruff check . --exclude data/raw --exclude .venv`.
- Look for brittle SQL filters, hidden source assumptions, unhandled DB failures, and duplicated query logic.

## Output expected

Return findings first, ordered by severity, with file/line references where possible. Then provide a short summary of strengths and a prioritized fix list before dashboard deployment.
"""


def _write_markdown(path: Path, content: str) -> None:
    full_path = PROJECT_ROOT / path
    full_path.parent.mkdir(parents=True, exist_ok=True)
    full_path.write_text(content, encoding="utf-8")


def _write_sheet(workbook: Workbook, title: str, rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title=title)
    for row in rows:
        sheet.append([_excel_value(value) for value in row])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width + 2, 42)


def _excel_value(value: object) -> object:
    """Convert values to types supported by openpyxl."""
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        if value.tzinfo is not None:
            return value.isoformat()
        return value.to_pydatetime()
    return value


def write_excel_workbook() -> Path:
    """Write an Excel workbook with executive and data quality tabs."""
    frames = load_dashboard_frames()
    summary = dashboard_summary(frames)

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(
        workbook,
        "Summary",
        [
            ["Metric", "Value"],
            ["Top BC vs Canada gap", f"{summary['top_gap_procedure']} ({summary['top_gap_days']:.1f} days)"],
            [
                "Top VCH wait",
                f"{summary['top_vch_wait_procedure']} ({summary['top_vch_wait_days']:.1f} days)",
            ],
            ["BC null case-volume rows", summary["bc_missing_case_volume_rows"]],
            ["Validity issues", summary["validity_issues"]],
        ],
    )
    for sheet_name, frame_name in [
        ("National Snapshot", "executive_national"),
        ("BC Gaps", "bc_gaps"),
        ("VCH Long Waits", "vch_long_waits"),
        ("Quality Missingness", "quality_missingness"),
        ("Freshness", "quality_freshness"),
    ]:
        frame = frames[frame_name]
        _write_sheet(workbook, sheet_name[:31], [list(frame.columns)] + frame.head(100).values.tolist())

    output_path = PROJECT_ROOT / EXCEL_FILE
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def build_deck_source() -> str:
    """Build the Markdown source text for the six-page executive deck."""
    frames = load_dashboard_frames()
    summary = dashboard_summary(frames)
    return f"""# Board Briefing Deck Source

## Slide 1: Executive Summary

- Largest BC vs Canada gap: {summary["top_gap_procedure"]} (+{summary["top_gap_days"]:.1f} days).
- Longest VCH median wait: {summary["top_vch_wait_procedure"]} ({summary["top_vch_wait_days"]:.1f} days).
- Main reporting caveat: BC_MoH null case-volume rows ({summary["bc_missing_case_volume_rows"]:,} rows).

## Slide 2: KPI Snapshot

- National snapshot: CT and MRI carry the largest latest-year case volumes.
- Orthopedic waits remain leadership-relevant: hip and knee replacement have high median and p90 waits.

## Slide 3: Trend and Variation

- Use 5-year national trends for high-volume procedures.
- Use BC vs Canada gaps to separate local variance from national pressure.

## Slide 4: Vancouver Coastal Operational View

- Prioritize VCH long-wait procedure drilldowns.
- Add hospital-level p90 review before operational claims.

## Slide 5: Data Quality and Reporting Risks

- Duplicate keys: {summary["duplicate_rows"]}.
- Numeric validity issues: {summary["validity_issues"]}.
- Suppressed/null volumes must remain visible.

## Slide 6: Recommended Next Steps

- Keep source unit assumptions visible in dashboard caveats.
- Review dashboard copy with healthcare-domain readers.
- Add deployment monitoring after Streamlit launch.
"""


def write_briefing_deck_pdf() -> Path:
    """Write a simple six-page PDF briefing deck using matplotlib."""
    source = build_deck_source()
    _write_markdown(DECK_SOURCE_FILE, source)
    slides = [block.strip() for block in source.split("## Slide ")[1:]]
    output_path = PROJECT_ROOT / DECK_PDF_FILE
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(output_path) as pdf:
        for slide in slides:
            title, *body = slide.splitlines()
            fig = plt.figure(figsize=(11, 8.5))
            fig.patch.set_facecolor("white")
            fig.text(0.08, 0.86, f"Slide {title}", fontsize=22, weight="bold", color="#17324D")
            y = 0.72
            for line in body:
                if not line.strip():
                    continue
                fig.text(0.1, y, line, fontsize=15, color="#263238")
                y -= 0.075
            fig.text(0.08, 0.08, "Public aggregate data demo. No patient-level or internal VCH data.", fontsize=10)
            pdf.savefig(fig)
            plt.close(fig)
    return output_path


def write_portfolio_artifacts() -> dict[str, Path]:
    """Write all final portfolio artifacts."""
    _write_markdown(METHODOLOGY_FILE, build_methodology_note())
    _write_markdown(REVIEW_GUIDE_FILE, build_review_guide())
    deck_path = write_briefing_deck_pdf()
    excel_path = write_excel_workbook()
    return {
        "methodology": PROJECT_ROOT / METHODOLOGY_FILE,
        "review_guide": PROJECT_ROOT / REVIEW_GUIDE_FILE,
        "deck_pdf": deck_path,
        "excel": excel_path,
    }


def main() -> None:
    """Generate portfolio artifacts and print written paths."""
    for name, path in write_portfolio_artifacts().items():
        print(f"{name}: {path.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
